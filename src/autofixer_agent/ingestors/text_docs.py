from __future__ import annotations

import json
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from autofixer_agent.rag.chromadb_store import RagDocument
from autofixer_agent.rag.chunking import chunk_text


def _read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rst", ".log", ".yaml", ".yml", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".json":
        return json.dumps(json.loads(path.read_text(encoding="utf-8", errors="ignore")), indent=2)
    if suffix in {".html", ".htm", ".xml"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return soup.get_text(separator="\n")
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Unsupported text document extension: {suffix}")


def ingest_text_document(path: str) -> list[RagDocument]:
    file_path = Path(path)
    text = _read_text(file_path)
    return [
        RagDocument(
            content=chunk,
            metadata={"source_type": "text", "path": str(path), "extension": file_path.suffix.lower()},
        )
        for chunk in chunk_text(text)
    ]
