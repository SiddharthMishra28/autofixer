from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from autofixer_agent.rag.chromadb_store import RagDocument
from autofixer_agent.rag.chunking import chunk_text


def ingest_excel(path: str) -> list[RagDocument]:
    workbook = load_workbook(Path(path), data_only=False)
    documents: list[RagDocument] = []

    for sheet in workbook.worksheets:
        lines: list[str] = []
        merged_cells = [str(rng) for rng in sheet.merged_cells.ranges]
        if merged_cells:
            lines.append(f"Merged ranges: {', '.join(merged_cells)}")

        for row in sheet.iter_rows(values_only=False):
            parts: list[str] = []
            for cell in row:
                if cell.value is None:
                    continue
                parts.append(f"{cell.coordinate}={cell.value}")
            if parts:
                lines.append(" | ".join(parts))

        sheet_text = "\n".join(lines)
        for chunk in chunk_text(sheet_text):
            documents.append(
                RagDocument(
                    content=chunk,
                    metadata={"source_type": "excel", "sheet": sheet.title, "path": str(path)},
                )
            )

    return documents
