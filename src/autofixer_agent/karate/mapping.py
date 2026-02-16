from __future__ import annotations

import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path



@dataclass
class MappingRecord:
    scenario_id: str
    json_path: str
    expected_value: str
    comparator: str
    priority: str
    source_sheet: str
    source_cell: str


def _canonical_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def _normalize_headers(headers: list[str]) -> dict[str, int]:
    indexed = {_canonical_key(header): i for i, header in enumerate(headers) if header}

    aliases = {
        "scenario_id": ["scenarioid", "scenario", "testcase", "testcaseid", "tcid"],
        "json_path": ["jsonpath", "path", "fieldpath", "keypath"],
        "expected_value": ["expected", "expectedvalue", "value", "expectedresult"],
        "comparator": ["comparator", "operator", "assertion", "matchtype"],
        "priority": ["priority", "severity", "rank"],
    }

    resolved: dict[str, int] = {}
    for canonical, options in aliases.items():
        for option in options:
            if option in indexed:
                resolved[canonical] = indexed[option]
                break
    return resolved


def normalize_excel_mapping(path: str) -> list[MappingRecord]:
    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), data_only=True)
    records: list[MappingRecord] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=False))
        if not rows:
            continue

        header_values = [str(cell.value).strip() if cell.value is not None else "" for cell in rows[0]]
        header_map = _normalize_headers(header_values)
        required = {"scenario_id", "json_path", "expected_value"}
        if not required.issubset(header_map.keys()):
            continue

        for row in rows[1:]:
            def _get(col_name: str, default: str = "") -> str:
                idx = header_map.get(col_name)
                if idx is None or idx >= len(row):
                    return default
                value = row[idx].value
                return "" if value is None else str(value).strip()

            scenario_id = _get("scenario_id")
            json_path = _get("json_path")
            expected_value = _get("expected_value")
            comparator = _get("comparator", "==") or "=="
            priority = _get("priority", "normal") or "normal"

            if not scenario_id or not json_path:
                continue

            source_idx = header_map["scenario_id"]
            source_cell = row[source_idx].coordinate if source_idx < len(row) else ""

            records.append(
                MappingRecord(
                    scenario_id=scenario_id,
                    json_path=json_path,
                    expected_value=expected_value,
                    comparator=comparator,
                    priority=priority,
                    source_sheet=sheet.title,
                    source_cell=source_cell,
                )
            )

    return records


def mapping_records_to_json(records: list[MappingRecord]) -> str:
    return json.dumps([asdict(record) for record in records], indent=2)
